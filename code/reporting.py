"""Build the rubric-required results.xlsx strictly from real run artifacts."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from eval import CLASS_NAMES, compute_metrics, read_pred  # noqa: E402


SHEETS = {
    "Backbones": ["exp_id", "backbone", "weight_tag", "params_m", "gmac", "img_size", "epochs",
                  "seed", "val_macro_f1", "val_top1", "train_seconds_per_epoch", "latency_batch1_ms",
                  "notes"],
    "Training": ["exp_id", "backbone", "axis", "difference_from_T00", "seed", "val_macro_f1",
                 "val_top1", "delta_vs_T00", "rare_class_f1", "notes"],
    "Inference": ["exp_id", "method", "checkpoint", "k", "val_macro_f1", "val_top1", "val_ece",
                  "p50_ms", "p95_ms", "p99_ms", "images_per_s", "relative_cost"],
    "Final": ["exp_id", "configuration", "seed", "val_macro_f1", "test_macro_f1", "test_top1",
              "test_ece", "mean_std"],
    "PerClass": ["exp_id", "seed", "class", "test_support", "precision", "recall", "f1"],
    "Latency": ["configuration", "gpu", "dtype", "batch", "img_size", "fused_bn", "p50_ms",
                "p95_ms", "p99_ms", "images_per_s", "torch"],
    "Summary": ["rank", "exp_id", "seed", "backbone", "val_macro_f1", "val_top1", "params_m",
                "gmac", "epoch_seconds", "test_macro_f1", "notes"],
}


def _load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def _run_records(runs_dir: Path) -> list[dict]:
    records = []
    for summary_path in sorted(runs_dir.glob("*/seed*/summary.json")):
        summary = _load_json(summary_path)
        config_path = summary_path.with_name("config.json")
        config = _load_json(config_path) if config_path.exists() else {}
        records.append({**config, **summary, "run_dir": str(summary_path.parent)})
    return records


def _optional_csv(path: Path, columns: list[str]) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=columns)
    frame = pd.read_csv(path)
    return frame.reindex(columns=columns)


def _axis_and_difference(config: dict) -> tuple[str, str]:
    init = config.get("init", "finetune")
    aug = config.get("aug", "basic")
    loss = config.get("loss", "ce")
    sampler = config.get("sampler")
    mix = config.get("mix")
    ema = config.get("ema_decay")
    if init != "finetune":
        return "A-init", f"init={init}"
    if aug != "basic":
        return "B-augmentation", f"aug={aug}"
    if loss != "ce":
        return "C-loss", f"loss={loss}"
    if sampler is not None:
        return "D-sampler", f"sampler={sampler}"
    if mix is not None:
        return "B-augmentation", f"mix={mix}, alpha={config.get('mix_alpha')}"
    if ema is not None:
        return "F-regularization", f"ema_decay={ema}"
    return "baseline", "T00 baseline recipe"


def _prediction_metrics(pred_dir: Path):
    scalar_rows, class_rows = [], []
    pattern = re.compile(r"^(?P<exp>.+)_seed(?P<seed>\d+)_test\.csv$")
    for path in sorted(pred_dir.glob("*_seed*_test.csv")):
        match = pattern.match(path.name)
        if not match:
            continue
        pred = read_pred(str(path))
        metrics = compute_metrics(pred.y_true, pred.y_pred, pred.probs)
        exp_id, seed = match.group("exp"), int(match.group("seed"))
        scalar_rows.append({"exp_id": exp_id, "seed": seed, **{
            key: metrics[key] for key in ("macro_f1", "top1", "ece")
        }})
        for index, class_name in enumerate(CLASS_NAMES):
            class_rows.append({
                "exp_id": exp_id, "seed": seed, "class": class_name,
                "test_support": int(metrics["support"][index]),
                "precision": float(metrics["precision"][index]),
                "recall": float(metrics["recall"][index]),
                "f1": float(metrics["f1"][index]),
            })
    return pd.DataFrame(scalar_rows), pd.DataFrame(class_rows, columns=SHEETS["PerClass"])


def _mean_std_text(values) -> str:
    array = np.asarray(values, dtype=float)
    if len(array) == 0:
        return ""
    if len(array) == 1:
        return f"{array[0]:.4f} (1 seed)"
    return f"{array.mean():.4f} ± {array.std(ddof=1):.4f}"


def build_results_workbook(output: str | Path = "results.xlsx", runs_dir: str | Path = "runs",
                           pred_dir: str | Path = "predictions") -> Path:
    """Create seven required sheets; absent experiments remain blank rather than fabricated."""
    runs_path, predictions_path = Path(runs_dir), Path(pred_dir)
    records = _run_records(runs_path)
    test_scalar, per_class = _prediction_metrics(predictions_path)

    backbone_rows, training_rows, final_rows, summary_rows = [], [], [], []
    baseline_by_seed = {
        int(row["seed"]): float(row["val_macro_f1"])
        for row in records if row.get("exp_id") == "T00" and "val_macro_f1" in row
    }
    test_lookup = {(row.exp_id, int(row.seed)): row for row in test_scalar.itertuples()}

    for record in records:
        exp_id = str(record.get("exp_id", ""))
        seed = int(record.get("seed", 0))
        common = {
            "exp_id": exp_id,
            "backbone": record.get("backbone"),
            "seed": seed,
            "val_macro_f1": record.get("val_macro_f1"),
            "val_top1": record.get("val_top1"),
        }
        if exp_id.startswith("B"):
            backbone_rows.append({
                **common, "weight_tag": record.get("weight_tag"), "params_m": record.get("params_m"),
                "gmac": record.get("gmac"), "img_size": record.get("img_size"),
                "epochs": record.get("epochs"), "train_seconds_per_epoch": record.get("epoch_seconds_mean"),
                "latency_batch1_ms": None, "notes": "",
            })
        if exp_id.startswith("T"):
            axis, difference = _axis_and_difference(record)
            base = baseline_by_seed.get(seed)
            value = record.get("val_macro_f1")
            training_rows.append({
                **common, "axis": axis, "difference_from_T00": difference,
                "delta_vs_T00": (float(value) - base if base is not None and value is not None else None),
                "rare_class_f1": None, "notes": "",
            })
        test = test_lookup.get((exp_id, seed))
        if exp_id.startswith("F") or test is not None:
            final_rows.append({
                "exp_id": exp_id,
                "configuration": f"{record.get('backbone')} | {record.get('loss')} | {record.get('aug')}",
                "seed": seed, "val_macro_f1": record.get("val_macro_f1"),
                "test_macro_f1": getattr(test, "macro_f1", None),
                "test_top1": getattr(test, "top1", None), "test_ece": getattr(test, "ece", None),
                "mean_std": "",
            })
        summary_rows.append({
            "exp_id": exp_id, "seed": seed, "backbone": record.get("backbone"),
            "val_macro_f1": record.get("val_macro_f1"), "val_top1": record.get("val_top1"),
            "params_m": record.get("params_m"), "gmac": record.get("gmac"),
            "epoch_seconds": record.get("epoch_seconds_mean"),
            "test_macro_f1": getattr(test, "macro_f1", None), "notes": "",
        })

    final = pd.DataFrame(final_rows, columns=SHEETS["Final"])
    if not final.empty:
        for exp_id, group in final.groupby("exp_id"):
            text = _mean_std_text(group["test_macro_f1"].dropna())
            final.loc[group.index[0], "mean_std"] = text
    summary = pd.DataFrame(summary_rows)
    if not summary.empty:
        summary = summary.sort_values("val_macro_f1", ascending=False, na_position="last").head(10).reset_index(drop=True)
        summary.insert(0, "rank", np.arange(1, len(summary) + 1))
    summary = summary.reindex(columns=SHEETS["Summary"])

    frames = {
        "Backbones": pd.DataFrame(backbone_rows, columns=SHEETS["Backbones"]),
        "Training": pd.DataFrame(training_rows, columns=SHEETS["Training"]),
        "Inference": _optional_csv(runs_path / "inference_results.csv", SHEETS["Inference"]),
        "Final": final,
        "PerClass": per_class,
        "Latency": _optional_csv(runs_path / "latency_results.csv", SHEETS["Latency"]),
        "Summary": summary,
    }

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.book[sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for column_index, column in enumerate(worksheet.columns, start=1):
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
                worksheet.column_dimensions[get_column_letter(column_index)].width = width
    return output_path
